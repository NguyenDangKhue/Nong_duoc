# -*- coding: utf-8 -*-
"""Tạo file Excel tổng hợp cây trồng × thuốc cấm."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

wb = Workbook()

# --- dữ liệu ---
GROUPS = {
    "hanh": "Hành, hẹ, gia vị",
    "la": "Rau ăn lá",
    "cu": "Rau củ",
    "qua": "Quả, đậu, khác",
}

AIS = [
    ("petroleum_oil", "Dầu khoáng", "SK ENSPRAY 99% EC", "Petroleum spray oil"),
    ("myclobutanil", "Myclobutanil", "Kanaka 405WP", ""),
    ("metalaxyl", "Metalaxyl", "METAXYL 25WP", ""),
    ("dimethomorph", "Dimethomorph", "CYLEN 500WP, INSURAN 50WP, Phytocide 50WP, EDDY 72WP", "Eddy cũng chứa dimethomorph"),
    ("boscalid", "Boscalid", "KIMONO APC 50WG", ""),
    ("imidacloprid", "Imidacloprid", "MAP JONO 700WP, Mikhada 10WP, Anvado 100WP", ""),
    ("clothianidin", "Clothianidin", "Dantotsu 50WG", ""),
    ("thiamethoxam", "Thiamethoxam", "Thiamax 25WG, Actara 25WG", ""),
    ("iprodione", "Iprodione", "Viroval 50WP / 50BTN", ""),
    ("permethrin", "Permethrin", "Map-Permethrin 50EC", ""),
    ("cyazofamid", "Cyazofamid", "Ranman 10SC", ""),
    ("lambda_cyhalothrin", "Cyhalothrin", "Icon 2.5EC", "Lambda-cyhalothrin"),
    ("azoxystrobin", "Azoxystrobin", "Amistar 250SC", ""),
    ("acetamiprid", "Acetamiprid", "Mospilan 3EC", ""),
    ("chlorfenapyr", "Chlorfenapyr", "SECURE 10EC", ""),
    ("novaluron", "Novaluron", "Rimon 10EC", ""),
    ("diazinon", "Diazinon", "DIAZAN 10H", "Hạt rải, không pha bình phun"),
    ("fosthiazate", "Fosthiazate", "Makeno 10GR", "Hạt rải, không pha bình phun"),
    ("chlorantraniliprole", "Chlorantraniliprole", "Dupont Prevathon 5SC, Dupont Prevathon 200SC", ""),
    ("cyantraniliprole", "Cyantraniliprole", "BENEVIA 100OD", ""),
    ("cypermethrin", "Cypermethrin", "Sec Saigon 10EC, NP-Cyrin Super", ""),
]

CROPS = [
    ("hanh-tay", "hanh", "Hành tây", "玉ねぎ", [
        "petroleum_oil", "myclobutanil", "metalaxyl", "dimethomorph",
        "imidacloprid", "clothianidin", "thiamethoxam", "iprodione",
        "chlorfenapyr", "novaluron", "diazinon", "fosthiazate",
        "chlorantraniliprole", "cyantraniliprole", "cypermethrin",
    ]),
    ("he", "hanh", "Hẹ", "ニラ", [
        "petroleum_oil", "dimethomorph", "permethrin", "imidacloprid",
        "chlorfenapyr", "novaluron", "diazinon", "fosthiazate",
    ]),
    ("hanh-huong", "hanh", "Hành hương", "わけぎ", [
        "petroleum_oil", "chlorantraniliprole", "cyantraniliprole",
    ]),
    ("hanh", "hanh", "Hành", "ネギ", ["dimethomorph", "boscalid", "cyazofamid", "novaluron"]),
    ("mitsuba", "hanh", "Rau cần ta (Mitsuba)", "みつば", [
        "myclobutanil", "metalaxyl", "dimethomorph", "permethrin",
        "thiamethoxam", "iprodione", "cyazofamid", "novaluron",
    ]),
    ("xa-lach", "la", "Xà lách", "レタス", [
        "myclobutanil", "permethrin", "lambda_cyhalothrin", "iprodione",
        "chlorfenapyr", "novaluron", "diazinon", "chlorantraniliprole", "cypermethrin",
    ]),
    ("xa-lach-xoong", "la", "Xà lách xoong", "クレソン", [
        "myclobutanil", "dimethomorph", "iprodione", "novaluron",
    ]),
    ("bo-xoi", "la", "Bó xôi", "ほうれん草", ["myclobutanil", "cyantraniliprole"]),
    ("cai-thia", "la", "Cải thìa", "チンゲン菜", [
        "petroleum_oil", "myclobutanil", "dimethomorph", "permethrin", "iprodione",
        "chlorfenapyr", "novaluron", "diazinon", "fosthiazate",
    ]),
    ("cai-thao", "la", "Cải thảo", "白菜", [
        "petroleum_oil", "myclobutanil", "iprodione",
    ]),
    ("bap-cai", "la", "Bắp cải", "キャベツ", [
        "petroleum_oil", "myclobutanil", "iprodione", "novaluron", "fosthiazate",
    ]),
    ("cai-ngong", "la", "Cải ngọt (Komatsuna)", "小松菜", ["novaluron"]),
    ("la-cu-cai", "la", "Lá củ cải", "葉大根", [
        "petroleum_oil", "myclobutanil", "iprodione", "diazinon", "fosthiazate",
    ]),
    ("rau-vn", "la", "Rau bán ở Việt Nam", "現地向け品種", []),
    ("ca-rot", "cu", "Cà rốt", "ニンジン", [
        "petroleum_oil", "myclobutanil", "dimethomorph", "chlorantraniliprole",
    ]),
    ("cu-cai", "cu", "Củ cải", "大根", [
        "petroleum_oil", "myclobutanil", "iprodione", "diazinon", "fosthiazate",
    ]),
    ("khoai-lang", "cu", "Khoai lang (Caiapo)", "カイアポイモ", [
        "petroleum_oil", "myclobutanil", "metalaxyl", "dimethomorph",
        "iprodione", "fosthiazate", "cyantraniliprole",
    ]),
    ("bi-ngo", "qua", "Bí ngô / Bí ngòi", "かぼちゃ", [
        "petroleum_oil", "thiamethoxam", "fosthiazate", "cyantraniliprole", "cypermethrin",
    ]),
    ("dau-nhat", "qua", "Đậu Nhật / Đậu cô ve", "インゲン", [
        "myclobutanil", "metalaxyl", "dimethomorph", "azoxystrobin",
        "iprodione", "cyazofamid", "novaluron", "fosthiazate",
    ]),
    ("rau-day", "qua", "Rau đay", "モロヘイヤ", [
        "lambda_cyhalothrin", "imidacloprid", "chlorfenapyr", "fosthiazate",
    ]),
    ("bap-ngot", "qua", "Bắp ngọt", "スイートコーン", ["diazinon", "fosthiazate"]),
    ("ca-tim", "qua", "Cà tím", "ナス", [
        "dimethomorph", "acetamiprid", "clothianidin", "thiamethoxam",
        "chlorfenapyr", "novaluron", "chlorantraniliprole",
    ]),
    ("gung", "qua", "Gừng", "生姜", []),
]

EMPTY_CROP_ROWS = 8

thin = Border(
    left=Side(style="thin", color="C5D0C8"),
    right=Side(style="thin", color="C5D0C8"),
    top=Side(style="thin", color="C5D0C8"),
    bottom=Side(style="thin", color="C5D0C8"),
)
thick_right = Border(
    left=Side(style="thin", color="C5D0C8"),
    right=Side(style="medium", color="1F6B45"),
    top=Side(style="thin", color="C5D0C8"),
    bottom=Side(style="thin", color="C5D0C8"),
)

fill_nav = PatternFill("solid", fgColor="0F2419")
fill_green = PatternFill("solid", fgColor="1F6B45")
fill_green2 = PatternFill("solid", fgColor="164D32")
fill_head2 = PatternFill("solid", fgColor="E7F3EB")
fill_ban = PatternFill("solid", fgColor="FDECEA")
fill_ok = PatternFill("solid", fgColor="FFFFFF")
fill_alt = PatternFill("solid", fgColor="F7F8F4")
fill_info = PatternFill("solid", fgColor="E8F0F8")
fill_warn = PatternFill("solid", fgColor="FEF3E2")
fill_empty = PatternFill("solid", fgColor="FFFBEB")
fill_group = {
    "hanh": PatternFill("solid", fgColor="F3EEE4"),
    "la": PatternFill("solid", fgColor="E6F4EF"),
    "cu": PatternFill("solid", fgColor="E8F0F8"),
    "qua": PatternFill("solid", fgColor="F5E9F7"),
}

font_white = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_white_sm = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
font_head = Font(name="Calibri", size=11, bold=True, color="14211A")
font_sm = Font(name="Calibri", size=8, color="5C6B62")
font_body = Font(name="Calibri", size=11, color="14211A")
font_ban = Font(name="Calibri", size=10, bold=True, color="B42318")
font_title = Font(name="Calibri", size=16, bold=True, color="0F2419")
font_note = Font(name="Calibri", size=10, color="5C6B62")

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

ai_map = {k: (name, prods, note) for k, name, prods, note in AIS}

# ========== Sheet 1: Hướng dẫn ==========
ws0 = wb.active
ws0.title = "Hướng dẫn"

ws0.merge_cells("A1:F1")
ws0["A1"] = "Sổ cây trồng × thuốc cấm phun"
ws0["A1"].font = font_title
ws0["A1"].alignment = left
ws0.row_dimensions[1].height = 28

ws0.merge_cells("A2:F2")
ws0["A2"] = "Dùng file này để bổ sung / sửa danh sách. Sau khi sửa, gửi lại file để cập nhật chương trình Tra pha."
ws0["A2"].font = font_note
ws0.merge_cells("A3:F3")
ws0["A3"] = "Nguồn: ô 使用不可 trên bảng chuẩn. Ô màu đỏ = CẤM. Ô trống = được dùng (trong phạm vi bảng)."
ws0["A3"].font = font_note

guide = [
    ("Sheet «Ma trận CẤM» (chỉnh ở đây cho nhanh)", [
        "Mỗi hàng = 1 loại cây. Mỗi cột (từ cột E trở đi) = 1 hoạt chất.",
        "Gõ CẤM vào ô nếu cây đó không được dùng hoạt chất đó. Xóa chữ CẤM nếu được dùng.",
        "Thêm cây mới: điền vào các hàng trống màu vàng ở cuối bảng (nhóm, tên Việt, tên Nhật).",
        "Thêm hoạt chất mới: chèn cột bên phải, ghi mã (vd. mancozeb) ở hàng 2, tên ở hàng 3, sản phẩm ở hàng 4, rồi đánh CẤM từng cây.",
        "Đừng đổi chữ ở hàng 2 (mã hoạt chất) trừ khi thêm cột mới — chương trình đọc theo mã này.",
    ]),
    ("Sheet «Danh sách chi tiết»", [
        "Mỗi hàng = một cặp cây + hoạt chất bị cấm. Tiện in hoặc lọc.",
        "Sheet này được tạo sẵn từ ma trận hiện tại. Nếu bạn sửa ma trận, hãy ưu tiên ma trận làm nguồn chính.",
    ]),
    ("Sheet «Cây trồng» / «Hoạt chất»", [
        "Danh mục gốc: mã cây, nhóm, tên Việt/Nhật; mã hoạt chất và sản phẩm tương ứng.",
        "Có thể ghi chú thêm ở cột Ghi chú.",
    ]),
    ("Quy ước", [
        "CẤM theo HOẠT CHẤT, không theo từng chai. Ví dụ cấm Dimethomorph thì cả Cylen, Insuran, Phytocide và Eddy đều bị khóa.",
        "DIAZAN 10H và Makeno 10GR là hạt rải gốc — vẫn ghi CẤM trên cây, nhưng không pha vào bình phun lá.",
        "Nhãn trên bao bì luôn thắng bảng này nếu khác nhau.",
    ]),
]

r = 5
for title, bullets in guide:
    ws0.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws0.cell(r, 1, title).font = font_white
    ws0.cell(r, 1).fill = fill_green
    ws0.cell(r, 1).alignment = left
    for c in range(1, 7):
        ws0.cell(r, c).fill = fill_green
    r += 1
    for b in bullets:
        ws0.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws0.cell(r, 1, "•  " + b).font = font_body
        ws0.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws0.row_dimensions[r].height = 32
        r += 1
    r += 1

ws0.column_dimensions["A"].width = 22
for col in "BCDEF":
    ws0.column_dimensions[col].width = 18
ws0.freeze_panes = "A5"
ws0.sheet_properties.tabColor = "1F6B45"

# ========== Sheet 2: Ma trận ==========
ws = wb.create_sheet("Ma trận CẤM", 0)
# put matrix first for convenience - actually user asked for excel to edit, matrix should be first
# I created Huong dan as active then created matrix at index 0, so matrix is first. Move huong dan? 
# create_sheet(..., 0) inserts at beginning, so Ma trận is first. Good.
# Wait I already had ws0 as active "Hướng dẫn", then create_sheet Ma trận at 0 - Ma trận becomes first.

info_cols = ["Nhóm", "Mã cây (đừng đổi)", "Tên cây (Việt)", "Tên Nhật"]
n_info = len(info_cols)
n_ai = len(AIS)
n_crop = len(CROPS)
last_data_row = 4 + n_crop + EMPTY_CROP_ROWS
last_col = n_info + n_ai

# Row 1 title
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
ws["A1"] = "MA TRẬN CẤM — gõ CẤM hoặc xóa. Hàng vàng = chỗ thêm cây mới. Chèn cột bên phải để thêm hoạt chất."
ws["A1"].font = font_white
ws["A1"].fill = fill_nav
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 22
for c in range(1, last_col + 1):
    ws.cell(1, c).fill = fill_nav
    ws.cell(1, c).font = font_white

# Row 2: mã AI
# Row 3: tên hoạt chất
# Row 4: sản phẩm
for c, h in enumerate(info_cols, 1):
    for row, fill, font in ((2, fill_green2, font_white_sm), (3, fill_green, font_white), (4, fill_head2, font_sm)):
        cell = ws.cell(row, c, h if row == 3 else ("Mã" if row == 2 and c == 2 else (h if row == 4 else "")))
        cell.fill = fill
        cell.font = font
        cell.alignment = center
        cell.border = thin

ws.cell(2, 1, "")
ws.cell(2, 2, "id")
ws.cell(2, 3, "")
ws.cell(2, 4, "")
for c in range(1, 5):
    ws.cell(2, c).fill = fill_green2
    ws.cell(2, c).font = font_white_sm
    ws.cell(2, c).alignment = center
    ws.cell(2, c).border = thick_right if c == 4 else thin
    ws.cell(3, c).fill = fill_green
    ws.cell(3, c).font = font_white
    ws.cell(3, c).alignment = center
    ws.cell(3, c).border = thick_right if c == 4 else thin
    ws.cell(4, c).fill = fill_head2
    ws.cell(4, c).font = font_sm
    ws.cell(4, c).alignment = center
    ws.cell(4, c).border = thick_right if c == 4 else thin

ws.cell(3, 1, "Nhóm")
ws.cell(3, 2, "Mã cây")
ws.cell(3, 3, "Tên cây (Việt)")
ws.cell(3, 4, "Tên Nhật")
ws.cell(4, 1, "Lọc theo nhóm")
ws.cell(4, 2, "Đừng đổi mã cũ")
ws.cell(4, 3, "Có thể sửa tên")
ws.cell(4, 4, "Có thể sửa")

for i, (key, name, prods, note) in enumerate(AIS):
    col = n_info + 1 + i
    c2 = ws.cell(2, col, key)
    c2.fill = fill_green2
    c2.font = font_white_sm
    c2.alignment = center
    c2.border = thin
    c3 = ws.cell(3, col, name)
    c3.fill = fill_green
    c3.font = font_white
    c3.alignment = center
    c3.border = thin
    c4 = ws.cell(4, col, prods)
    c4.fill = fill_head2
    c4.font = font_sm
    c4.alignment = center
    c4.border = thin
    if note:
        c3.comment = Comment(note, "Sổ BVTV")

ws.row_dimensions[2].height = 18
ws.row_dimensions[3].height = 32
ws.row_dimensions[4].height = 48

# data rows
dv = DataValidation(type="list", formula1='"CẤM"', allow_blank=True, showDropDown=False)
dv.prompt = "Chọn CẤM hoặc để trống nếu được phun"
dv.promptTitle = "Cấm phun?"
dv.error = "Chỉ nhập CẤM hoặc để trống"
dv.errorTitle = "Giá trị không hợp lệ"
ws.add_data_validation(dv)

for ri, (cid, gid, vn, ja, bans) in enumerate(CROPS):
    row = 5 + ri
    vals = [GROUPS[gid], cid, vn, ja]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row, c, v)
        cell.font = font_body
        cell.alignment = left if c >= 3 else center
        cell.border = thick_right if c == 4 else thin
        cell.fill = fill_group[gid]
    ban_set = set(bans)
    for i, (key, name, prods, note) in enumerate(AIS):
        col = n_info + 1 + i
        banned = key in ban_set
        cell = ws.cell(row, col, "CẤM" if banned else "")
        cell.alignment = center
        cell.border = thin
        cell.font = font_ban if banned else font_body
        cell.fill = fill_ban if banned else (fill_alt if ri % 2 else fill_ok)
    ws.row_dimensions[row].height = 20

# empty rows for new crops
for k in range(EMPTY_CROP_ROWS):
    row = 5 + n_crop + k
    for c in range(1, last_col + 1):
        cell = ws.cell(row, c, "")
        cell.border = thick_right if c == 4 else thin
        cell.fill = fill_empty
        cell.alignment = center
    ws.row_dimensions[row].height = 20

first_ai_col = n_info + 1
last_ai_col = last_col
dv.add(f"{get_column_letter(first_ai_col)}5:{get_column_letter(last_ai_col)}{last_data_row}")

# conditional formatting: CẤM -> red
red_font = Font(name="Calibri", size=10, bold=True, color="B42318")
ws.conditional_formatting.add(
    f"{get_column_letter(first_ai_col)}5:{get_column_letter(last_ai_col)}{last_data_row}",
    CellIsRule(operator="equal", formula=['"CẤM"'], fill=fill_ban, font=red_font),
)

# widths
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 26
ws.column_dimensions["D"].width = 16
for i in range(n_ai):
    ws.column_dimensions[get_column_letter(first_ai_col + i)].width = 14

ws.freeze_panes = "E5"
ws.auto_filter.ref = f"A4:{get_column_letter(last_col)}{last_data_row}"
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "B42318"
# freeze below headers: row 5 is first data, freeze E5 keeps A-D and rows 1-4
ws.sheet_view.zoomScale = 90
ws.print_title_rows = "1:4"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.oddHeader.left.text = "Cây × thuốc cấm"
ws.oddFooter.right.text = "Trang &P / &N"

# ========== Sheet 3: Danh sách chi tiết ==========
ws2 = wb.create_sheet("Danh sách chi tiết")
headers2 = [
    "Nhóm",
    "Mã cây",
    "Tên cây (Việt)",
    "Tên Nhật",
    "Mã hoạt chất",
    "Hoạt chất",
    "Sản phẩm bị cấm (cùng hoạt chất)",
    "Ghi chú",
]
for c, h in enumerate(headers2, 1):
    cell = ws2.cell(1, c, h)
    cell.font = font_white
    cell.fill = fill_green
    cell.alignment = center
    cell.border = thin
ws2.row_dimensions[1].height = 24

r = 2
for cid, gid, vn, ja, bans in CROPS:
    for key in bans:
        name, prods, note = ai_map[key]
        vals = [GROUPS[gid], cid, vn, ja, key, name, prods, note]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(r, c, v)
            cell.font = font_body
            cell.alignment = left
            cell.border = thin
            if r % 2 == 0:
                cell.fill = fill_alt
            if c == 6:
                cell.font = font_ban
        r += 1

# blank rows to add
for k in range(15):
    for c in range(1, 9):
        cell = ws2.cell(r + k, c, "")
        cell.border = thin
        cell.fill = fill_empty

ws2.auto_filter.ref = f"A1:H{r - 1}"
ws2.freeze_panes = "A2"
ws2.sheet_properties.tabColor = "B45309"
widths2 = [22, 16, 26, 16, 22, 22, 55, 36]
for i, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

# ========== Sheet 4: Cây trồng ==========
ws3 = wb.create_sheet("Cây trồng")
h3 = ["Mã cây", "Nhóm (mã)", "Nhóm (tên)", "Tên cây (Việt)", "Tên Nhật", "Số hoạt chất cấm", "Ghi chú / tên gọi khác"]
for c, h in enumerate(h3, 1):
    cell = ws3.cell(1, c, h)
    cell.font = font_white
    cell.fill = fill_green
    cell.alignment = center
    cell.border = thin
for ri, (cid, gid, vn, ja, bans) in enumerate(CROPS, 2):
    vals = [cid, gid, GROUPS[gid], vn, ja, len(bans), ""]
    for c, v in enumerate(vals, 1):
        cell = ws3.cell(ri, c, v)
        cell.font = font_body
        cell.alignment = left if c >= 4 else center
        cell.border = thin
        cell.fill = fill_group[gid]
for k in range(8):
    row = 2 + n_crop + k
    for c in range(1, 8):
        cell = ws3.cell(row, c, "")
        cell.border = thin
        cell.fill = fill_empty
ws3.auto_filter.ref = f"A1:G{1 + n_crop}"
ws3.freeze_panes = "A2"
for i, w in enumerate([16, 12, 22, 26, 16, 18, 40], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.sheet_properties.tabColor = "1F6B45"

# ========== Sheet 5: Hoạt chất ==========
ws4 = wb.create_sheet("Hoạt chất")
h4 = ["Mã hoạt chất", "Tên hoạt chất", "Sản phẩm trong sổ / cùng hoạt chất", "Ghi chú", "Số cây bị cấm"]
for c, h in enumerate(h4, 1):
    cell = ws4.cell(1, c, h)
    cell.font = font_white
    cell.fill = fill_green
    cell.alignment = center
    cell.border = thin
ban_count = {k: 0 for k, *_ in AIS}
for *_, bans in CROPS:
    for k in bans:
        ban_count[k] += 1
for ri, (key, name, prods, note) in enumerate(AIS, 2):
    vals = [key, name, prods, note, ban_count[key]]
    for c, v in enumerate(vals, 1):
        cell = ws4.cell(ri, c, v)
        cell.font = font_body
        cell.alignment = left
        cell.border = thin
        if ri % 2 == 0:
            cell.fill = fill_alt
for k in range(8):
    row = 2 + n_ai + k
    for c in range(1, 6):
        cell = ws4.cell(row, c, "")
        cell.border = thin
        cell.fill = fill_empty
ws4.auto_filter.ref = f"A1:E{1 + n_ai}"
ws4.freeze_panes = "A2"
for i, w in enumerate([24, 22, 62, 40, 16], 1):
    ws4.column_dimensions[get_column_letter(i)].width = w
ws4.sheet_properties.tabColor = "1D4E89"

# ========== Sheet 6: Thuốc chưa gắn cấm ==========
ws5 = wb.create_sheet("Thuốc chưa gắn cấm")
ws5["A1"] = "Các thuốc đang có trong sổ nhưng chưa bị cấm trên cây nào (theo bảng hiện tại)"
ws5["A1"].font = font_white
ws5["A1"].fill = fill_nav
ws5.merge_cells("A1:C1")
for c in range(1, 4):
    ws5.cell(1, c).fill = fill_nav
h5 = ["Tên thuốc", "Hoạt chất", "Ghi chú"]
for c, h in enumerate(h5, 1):
    cell = ws5.cell(2, c, h)
    cell.font = font_white
    cell.fill = fill_green
    cell.alignment = center
unlisted = [
    ("Dipel 6.4WG", "Bacillus thuringiensis (Bt)", "Chưa có cây cấm trong bảng này"),
    ("COC 85WP", "Copper oxychloride", "Khuyến cáo phun riêng, không phải cấm theo cây"),
    ("Daconil 500SC", "Chlorothalonil", "Chưa có cây cấm trong bảng này"),
    ("Neem Nim 0.3EC", "Azadirachtin", "Chưa có cây cấm trong bảng này"),
    ("Acti No Vate 1SP", "Streptomyces lydicus", "Chưa có cây cấm trong bảng này"),
]
for ri, rowv in enumerate(unlisted, 3):
    for c, v in enumerate(rowv, 1):
        cell = ws5.cell(ri, c, v)
        cell.font = font_body
        cell.border = thin
        cell.alignment = left
ws5.column_dimensions["A"].width = 24
ws5.column_dimensions["B"].width = 32
ws5.column_dimensions["C"].width = 48
ws5.sheet_properties.tabColor = "5C6B62"

# Reorder sheets: Ma trận, Hướng dẫn, Danh sách, Cây, Hoạt chất, Chưa gắn
# Current order after insert at 0: Ma trận, Hướng dẫn, Danh sách, Cây, Hoạt chất, Chưa gắn
# That's good. Move Hướng dẫn to index 1 (already). First is Ma trận.

out = r"d:\Heo Heo\Thông tin nông dược\Cay-trong-va-thuoc-cam-chuan.xlsx"
wb.save(out)
print("Wrote", out)
print("crops", n_crop, "ais", n_ai, "ban rows", sum(len(b[4]) for b in CROPS))
